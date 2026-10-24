from pathlib import Path
from openpyxl import load_workbook
from zipfile import ZipFile
import json,hashlib
import argparse
parser=argparse.ArgumentParser(description='Read-only independent verification of actual browser-downloaded Pushout XLSX files (openpyxl).')
parser.add_argument('--directory',default='output/pushout-95')
args=parser.parse_args()
root=Path(args.directory);result=[]
for f in sorted(root.glob('*.xlsx')):
 with ZipFile(f) as z: assert z.testzip() is None
 w=load_workbook(f,data_only=False);assert w.sheetnames==['Pushout','Metadata','Data notes']
 s=w['Pushout'];assert [s.cell(r,1).value for r in range(6,15)]==['Pot percentage','Missed pots','Fouls','Successful breaks','Dry/soft breaks','Foul breaks','Break and runs','Break and win','Break and lose']
 for row in s.iter_rows(min_row=6,max_row=14,min_col=2,max_col=3):
  for c in row:assert c.data_type=='n' and isinstance(c.value,(int,float)),(f,c.coordinate,c.value)
 assert s['B6'].number_format=='0%' and s['C6'].number_format=='0%'
 assert s['A3'].value=='This match only'
 for sh in w:
  for row in sh:
   assert not any(c.data_type=='f' for c in row)
 main=' '.join(str(c.value) for row in s for c in row)
 assert '*' not in main and 'a9f06904' not in main and 'unknown_attempts' not in main
 if f.name.startswith('singles'):
  assert list(s.values)[4]==('Metric','Daniel Lang','Tom Kurz')
  assert [[s.cell(r,c).value for c in [2,3]] for r in range(6,15)]==[[.93,.95],[1,2],[1,2],[1,3],[2,0],[0,1],[1,1],[1,3],[2,1]]
  assert dict(w['Metadata'].values)['Match ID']=='a9f06904-00a5-4518-a81a-eae5a9425a35'
 if f.name.startswith('team'):assert s['B5'].value=='Slovakia' and s['C5'].value=='Czechia'
 assert any('push out or soft break' in str(c.value) for row in w['Data notes'] for c in row)
 result.append(dict(file=f.name,typed=True,sheets=w.sheetnames,sha256=hashlib.sha256(f.read_bytes()).hexdigest()))
assert len(result)==6
Path('docs/pushout-95-xlsx-verification.json').write_text(json.dumps(result,indent=2))
print('PASS independent openpyxl + ZIP validation of six actual downloaded XLSX files: exact Singles values, numeric types/formats, notes, owners, no formulas/IDs in main')
